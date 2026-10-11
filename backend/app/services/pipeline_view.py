"""Read-only projection of Daria's canonical workbook for the test cockpit.

This module never writes the XLSX or persists a second application tracker.
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

SHEET = "Networking Pipeline"
AUDIT = "Metrics Audit"
HEADERS = (
    "Контакт", "Связь / контекст", "Компания / направление", "Что отправлено",
    "Ответ / реакция", "Следующий шаг", "Рынок", "Потолок зп",
    "Тех. база / заметки", "Дата первого контакта", "Дата последнего контакта",
    "Дней без ответа", "ЗП в EUR эквиваленте (год)",
    "Рыночная вилка (EUR/год) + источник", "Fit по компенсации", "SLA статус",
    "Employer score", "Соцслушание / сигнал", "Ключевой риск", "Источники",
    "Decision state", "Why this color", "Next trigger", "Next action", "Due",
    "Evidence to lead with", "Expected contact / review date", "Process ID",
)
URL_RE = re.compile(r"https?://[^\s;<>]+")
CLOSED = ("REJECT", "CLOSED", "NO FIT", "ARCHIVE")
REJECTED_SLA = ("REJECTED", "CANCELLED", "CLOSED", "REJECT —")


def _display(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else str(value)


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 30000 < value < 90000:
        return from_excel(value).date()
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            pass
    return None


def _folder_index(root):
    if not root.is_dir():
        return {}
    def key(name):
        return re.sub(r"\s+", " ", name.replace("✅", "").strip()).casefold()
    found = {}
    for child in root.iterdir():
        if child.is_dir() and child.name != "архив":
            found.setdefault(key(child.name), []).append(child)
    return found


def _lane(row, audit, today):
    state = _display(row[20]).upper().strip()
    sla = _display(row[15]).upper().strip()
    sent = _display(audit[3]).upper().strip() == "YES" or sla.startswith("APPLIED")
    conflict = any(state.startswith(x) for x in ("ACTION REQUIRED", "ACTIVE", "VERIFY FIRST")) and any(
        sla.startswith(x) for x in REJECTED_SLA
    )
    due = _date(row[24])
    # Old Due dates on waiting/monitor rows need reconciliation, not an invented task.
    conflict = conflict or bool(due and due < today and state.startswith(("DONE / WAITING", "WAITING", "MONITOR")))
    if any(state.startswith(x) for x in CLOSED):
        lane = "closed"
    elif conflict:
        lane = "verify"
    elif state.startswith("VERIFY FIRST"):
        lane = "verify"
    elif state.startswith(("ACTION REQUIRED", "ACTIVE")):
        lane = "do"
    elif state.startswith("MONITOR"):
        lane = "monitor"
    else:
        lane = "waiting"
    return lane, sent, conflict, due


def project(path: Path, today: date | None = None):
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    today = today or datetime.now(ZoneInfo("Europe/Madrid")).date()
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if set(book.sheetnames) != {SHEET, AUDIT}:
            raise ValueError("Unexpected workbook sheets")
        sheet, audit_sheet = book[SHEET], book[AUDIT]
        headers = tuple(next(sheet.iter_rows(min_row=13, max_row=13, max_col=28, values_only=True)))
        if headers != HEADERS:
            raise ValueError("Canonical A:AB headers do not match")
        audit = {}
        for raw in audit_sheet.iter_rows(min_row=2, values_only=True):
            if raw[0]:
                if raw[0] in audit:
                    raise ValueError(f"Duplicate Metrics Audit Process ID: {raw[0]}")
                audit[raw[0]] = raw
        records, ids = [], set()
        folders = _folder_index(path.parent)
        for raw in sheet.iter_rows(min_row=14, max_col=28, values_only=True):
            if not any(v is not None for v in raw):
                continue
            pid = raw[27]
            if not pid or pid in ids or pid not in audit:
                raise ValueError(f"Missing, duplicate or unjoined Process ID: {pid}")
            ids.add(pid)
            kind = _display(audit[pid][2])
            if kind != "Opportunity / process":
                continue
            lane, sent, conflict, due = _lane(raw, audit[pid], today)
            if lane == "closed":
                continue
            title = _display(raw[0])
            company, _, role = title.partition(" / ")
            if not role and " — " in title:
                company, _, role = title.partition(" — ")
            if not role and lane == "do":
                # An actionable unscoped contact is a verification task, not an application.
                lane, conflict = "verify", True
            urls = [u.rstrip(".,)") for u in URL_RE.findall(_display(raw[19]))]
            folder_key = re.sub(r"\s+", " ", title.replace(" / ", " - ").strip()).casefold()
            matches = folders.get(folder_key, [])
            records.append({
                "id": pid, "company": company, "role": role or "Роль не зафиксирована", "role_known": bool(role),
                "lane": lane, "sent": sent, "due": due.isoformat() if due else None,
                "due_text": _display(raw[24]), "next_action": _display(raw[23]) or _display(raw[5]),
                "blocker": _display(raw[18]), "state": _display(raw[20]), "sla": _display(raw[15]),
                "conflict": conflict, "overdue": bool(due and due < today),
                "sources": urls, "folder": str(matches[0]) if len(matches) == 1 else None,
                "fields": [{"label": h, "value": _display(v)} for h, v in zip(HEADERS, raw)],
            })
        if ids != set(audit):
            raise ValueError("Process ID set differs from Metrics Audit")
    finally:
        book.close()
    priority = {"do": 0, "verify": 1, "waiting": 2, "monitor": 3}
    records.sort(key=lambda r: (priority[r["lane"]], r["due"] is None, r["due"] or "9999", r["company"].casefold()))
    return {
        "source": path.name, "updated_at": datetime.fromtimestamp(path.stat().st_mtime, ZoneInfo("Europe/Madrid")).isoformat(timespec="minutes"),
        "read_only": True, "total": len(records),
        "counts": {lane: sum(r["lane"] == lane for r in records) for lane in priority},
        "records": records,
    }
