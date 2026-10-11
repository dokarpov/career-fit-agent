from __future__ import annotations
import csv, hashlib, json, re
from pathlib import Path
from bs4 import BeautifulSoup

TEXT_EXT = {".md",".txt",".html",".htm",".json",".csv",".docx",".xlsx"}

def stable_id(source_ref: str, title: str) -> str:
    h = hashlib.sha1(f"{source_ref}|{title}".encode()).hexdigest()[:12]
    return f"ev_{h}"

def read_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".md",".txt"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext in {".html",".htm"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
        return soup.get_text("\n", strip=True)
    if ext == ".json":
        obj = json.loads(path.read_text(encoding="utf-8"))
        return json.dumps(obj, ensure_ascii=False, indent=2)
    if ext == ".csv":
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v not in (None,"")]
                if vals:
                    parts.append(" | ".join(vals))
        return "\n".join(parts)
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    raise ValueError(f"Unsupported file type: {ext}")

def atomize_text(text: str, source_ref: str, source_type: str = "file") -> list[dict]:
    # Reference atomizer: headings / paragraphs become reviewable evidence candidates.
    chunks = [c.strip() for c in re.split(r"\n\s*\n|(?=^#{1,3}\s)", text, flags=re.M) if c.strip()]
    out = []
    for idx, chunk in enumerate(chunks[:500]):
        if len(chunk) < 40:
            continue
        title = re.sub(r"^#+\s*", "", chunk.splitlines()[0])[:120]
        evidence = chunk[:1800]
        out.append({
            "id": stable_id(source_ref, f"{idx}:{title}"),
            "source_type": source_type,
            "source_ref": source_ref,
            "title": title or f"Evidence fragment {idx+1}",
            "claim": title or evidence[:120],
            "evidence": evidence,
            "tags": [],
            "status": "needs-review",
            "direct": False,
            "visibility": "private",
            "boundary": "Imported candidate evidence; verify ownership and claim boundary before use."
        })
    return out

def ingest_path(path: str | Path) -> list[dict]:
    p = Path(path)
    if p.is_file():
        return atomize_text(read_text(p), str(p), "file")
    if p.is_dir():
        out = []
        for f in p.rglob("*"):
            if f.is_file() and (f.suffix.lower() in TEXT_EXT or f.suffix.lower()==".pdf"):
                try:
                    out.extend(atomize_text(read_text(f), str(f), "file"))
                except Exception:
                    pass
        return out
    raise FileNotFoundError(path)
