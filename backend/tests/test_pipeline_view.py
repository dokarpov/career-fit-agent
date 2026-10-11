from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from app.services.pipeline_view import HEADERS, project


class PipelineViewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "canonical.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Networking Pipeline"
        sheet.append([])
        for col, label in enumerate(HEADERS, 1):
            sheet.cell(13, col, label)
        audit = book.create_sheet("Metrics Audit")
        audit.append(["Process ID", "Record", "Record type", "Application evidence"])

        def add(rownum, pid, name, state, sla, due, evidence="No", kind="Opportunity / process"):
            values = [None] * 28
            values[0], values[15], values[20], values[24], values[27] = name, sla, state, due, pid
            for col, value in enumerate(values, 1):
                sheet.cell(rownum, col, value)
            audit.append([pid, name, kind, evidence])

        add(14, "P-1", "Alpha / PM", "ACTION REQUIRED", "Ready", datetime(2026, 10, 4))
        add(15, "P-2", "Beta / PM", "DONE / WAITING", "APPLIED — receipt", datetime(2026, 9, 1), "Yes")
        add(16, "P-3", "Gamma / PM", "REJECT", "REJECTED", None)
        add(17, "P-4", "Person", "ACTION REQUIRED", "New lead", None, kind="Relationship / network")
        add(18, "P-5", "Delta / PM", "MONITOR", "Watching", None)
        add(19, "P-6", "Lucia", "ACTION REQUIRED", "New lead", None)
        book.save(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_read_only_lanes_and_full_fields(self):
        before = self.path.read_bytes()
        result = project(self.path, date(2026, 10, 5))
        self.assertEqual(result["counts"], {"do": 1, "verify": 2, "waiting": 0, "monitor": 1})
        self.assertEqual([r["id"] for r in result["records"]], ["P-1", "P-2", "P-6", "P-5"])
        self.assertTrue(result["records"][1]["sent"])
        self.assertTrue(result["records"][1]["conflict"])
        self.assertFalse(result["records"][2]["role_known"])
        self.assertEqual(len(result["records"][0]["fields"]), 28)
        self.assertEqual(self.path.read_bytes(), before)

    def test_broken_join_fails_closed(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Networking Pipeline"
        for col, label in enumerate(HEADERS, 1):
            sheet.cell(13, col, label)
        sheet.cell(14, 1, "Alpha / PM")
        sheet.cell(14, 28, "MISSING")
        book.create_sheet("Metrics Audit")
        book.save(self.path)
        with self.assertRaises(ValueError):
            project(self.path)


if __name__ == "__main__":
    unittest.main()
